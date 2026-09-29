import io
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from server import db, get_current_user, require_module, hash_password, verify_password, log_audit, ALL_MODULES
from extra_routes import account_statement, load_settings

router = APIRouter()

# ==================== CHANGE OWN PASSWORD ====================

class PasswordChange(BaseModel):
    current_password: str
    new_password: str


@router.put("/auth/change-password")
async def change_password(data: PasswordChange, user: dict = Depends(get_current_user)):
    if len(data.new_password) < 6:
        raise HTTPException(status_code=400, detail="Le nouveau mot de passe doit contenir au moins 6 caractères")
    full = await db.users.find_one({"id": user["id"]})
    if not full or not verify_password(data.current_password, full["password"]):
        raise HTTPException(status_code=400, detail="Mot de passe actuel incorrect")
    await db.users.update_one({"id": user["id"]}, {"$set": {"password": hash_password(data.new_password)}})
    await log_audit(user, "update", "users", "users", user["id"], user["name"], after={"password": "(modifié par l'utilisateur)"})
    return {"message": "Mot de passe modifié"}


# ==================== CASH SESSION DETAIL (TICKET) ====================

@router.get("/cash/sessions/{session_id}")
async def get_session_detail(session_id: str, user: dict = Depends(require_module("cash"))):
    session = await db.cash_sessions.find_one({"id": session_id}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Session non trouvée")
    movements = await db.cash_movements.find({"session_id": session_id}, {"_id": 0}).sort("created_at", 1).to_list(10000)
    total_in = round(sum(m["amount"] for m in movements if m["type"] == "in"), 2)
    total_out = round(sum(m["amount"] for m in movements if m["type"] == "out"), 2)
    by_category = {}
    for m in movements:
        key = f"{m['type']}:{m.get('category', 'autre')}"
        by_category[key] = round(by_category.get(key, 0) + m["amount"], 2)
    expected = round(session.get("opening_amount", 0) + total_in - total_out, 2)
    return {
        "session": session,
        "movements": movements,
        "summary": {"total_in": total_in, "total_out": total_out, "expected": session.get("expected_amount") if session.get("expected_amount") is not None else expected, "by_category": by_category},
    }


# ==================== OVERDUE CUSTOMERS ====================

@router.get("/dashboard/overdue")
async def overdue_customers(days: int = 7, user: dict = Depends(get_current_user)):
    now = datetime.now(timezone.utc)
    sales = await db.sales.find({"balance_due": {"$gt": 0}, "status": {"$ne": "cancelled"}, "customer_id": {"$ne": None}}, {"_id": 0}).to_list(10000)
    grouped = {}
    for s in sales:
        created = s["created_at"] if isinstance(s["created_at"], datetime) else datetime.fromisoformat(s["created_at"])
        if created.tzinfo is None:
            created = created.replace(tzinfo=timezone.utc)
        late = (now - created).days
        if late < days:
            continue
        g = grouped.setdefault(s["customer_id"], {"customer_id": s["customer_id"], "customer_name": s.get("customer_name", ""), "overdue_amount": 0, "invoices": [], "days_late": 0})
        g["overdue_amount"] = round(g["overdue_amount"] + s["balance_due"], 2)
        g["days_late"] = max(g["days_late"], late)
        g["invoices"].append({"invoice_number": s["invoice_number"], "balance_due": s["balance_due"], "date": created.isoformat(), "days_late": late})
    if not grouped:
        return {"days": days, "customers": [], "total_overdue": 0}
    customers = await db.customers.find({"id": {"$in": list(grouped.keys())}}, {"_id": 0, "id": 1, "name": 1, "phone": 1, "balance": 1}).to_list(1000)
    cmap = {c["id"]: c for c in customers}
    result = []
    for cid, g in grouped.items():
        c = cmap.get(cid, {})
        g["customer_name"] = c.get("name", g["customer_name"])
        g["phone"] = c.get("phone", "")
        g["balance"] = c.get("balance", 0)
        g["invoices"].sort(key=lambda i: i["date"])
        result.append(g)
    result.sort(key=lambda g: -g["overdue_amount"])
    return {"days": days, "customers": result, "total_overdue": round(sum(g["overdue_amount"] for g in result), 2)}


# ==================== DAILY REPORT ====================

@router.get("/reports/daily")
async def daily_report(date: Optional[str] = None, user: dict = Depends(require_module("reports"))):
    day = date or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    rng = {"$gte": f"{day}T00:00:00", "$lte": f"{day}T23:59:59.999999"}

    sales = await db.sales.find({"created_at": rng}, {"_id": 0}).to_list(10000)
    active = [s for s in sales if s.get("status") != "cancelled"]
    cancelled = [s for s in sales if s.get("status") == "cancelled"]
    payments = await db.payments.find({"created_at": rng}, {"_id": 0}).to_list(10000)
    sale_payments = [p for p in payments if p.get("entity_type") == "sale"]
    customer_payments = [p for p in payments if p.get("entity_type") == "customer"]
    supplier_payments = [p for p in payments if p.get("entity_type") == "supplier"]

    by_method = {}
    def add(method, key, amount):
        m = by_method.setdefault(method or "especes", {"sales_cash": 0, "invoice_payments": 0, "customer_payments": 0, "supplier_payments": 0, "total_in": 0, "total_out": 0})
        m[key] = round(m[key] + amount, 2)
        if key == "supplier_payments":
            m["total_out"] = round(m["total_out"] + amount, 2)
        else:
            m["total_in"] = round(m["total_in"] + amount, 2)

    for s in active:
        later = sum(p["amount"] for p in sale_payments if p.get("entity_id") == s["id"])
        initial = round(s.get("amount_paid", 0) - later, 2)
        if initial > 0:
            add(s.get("payment_method"), "sales_cash", min(initial, s.get("total", initial)))
    for p in sale_payments:
        add(p.get("payment_method"), "invoice_payments", p["amount"])
    for p in customer_payments:
        add(p.get("payment_method"), "customer_payments", p["amount"])
    for p in supplier_payments:
        add(p.get("payment_method"), "supplier_payments", p["amount"])

    movements = await db.cash_movements.find({"created_at": rng}, {"_id": 0}).sort("created_at", 1).to_list(10000)
    cash_in = round(sum(m["amount"] for m in movements if m["type"] == "in"), 2)
    cash_out = round(sum(m["amount"] for m in movements if m["type"] == "out"), 2)
    cash_by_cat = {}
    for m in movements:
        k = f"{m['type']}:{m.get('category', 'autre')}"
        cash_by_cat[k] = round(cash_by_cat.get(k, 0) + m["amount"], 2)
    sessions = await db.cash_sessions.find({"$or": [{"opened_at": rng}, {"closed_at": rng}]}, {"_id": 0}).sort("opened_at", 1).to_list(100)

    bank = await db.bank_journal.find({"created_at": rng}, {"_id": 0}).sort("created_at", 1).to_list(10000)
    bank_credit = round(sum(b["amount"] for b in bank if b.get("entry_type") == "credit"), 2)
    bank_debit = round(sum(b["amount"] for b in bank if b.get("entry_type") == "debit"), 2)

    products_sold = {}
    for s in active:
        for it in s.get("items", []):
            pr = products_sold.setdefault(it["product_id"], {"product_name": it["product_name"], "quantity": 0, "total": 0})
            pr["quantity"] = round(pr["quantity"] + it["quantity"], 2)
            pr["total"] = round(pr["total"] + it["total"], 2)
    top_products = sorted(products_sold.values(), key=lambda x: -x["total"])[:10]

    total_revenue = round(sum(s.get("total", 0) for s in active), 2)
    total_received_at_sale = round(sum(min(s.get("amount_paid", 0), s.get("total", 0)) for s in active), 2)
    return {
        "date": day,
        "sales": {
            "count": len(active),
            "cancelled_count": len(cancelled),
            "total_revenue": total_revenue,
            "total_discount": round(sum(s.get("discount", 0) for s in active), 2),
            "total_credit": round(sum(s.get("balance_due", 0) for s in active), 2),
            "total_received": total_received_at_sale,
            "list": [{"invoice_number": s["invoice_number"], "customer_name": s.get("customer_name") or "Client Comptoir", "total": s.get("total", 0), "amount_paid": s.get("amount_paid", 0), "balance_due": s.get("balance_due", 0), "payment_method": s.get("payment_method"), "status": s.get("status"), "cashier_name": s.get("cashier_name", ""), "created_at": s["created_at"]} for s in sorted(sales, key=lambda x: x["created_at"])],
        },
        "by_method": by_method,
        "encaissements": {
            "invoice_payments": round(sum(p["amount"] for p in sale_payments), 2),
            "customer_payments": round(sum(p["amount"] for p in customer_payments), 2),
            "supplier_payments": round(sum(p["amount"] for p in supplier_payments), 2),
        },
        "cash": {"total_in": cash_in, "total_out": cash_out, "net": round(cash_in - cash_out, 2), "by_category": cash_by_cat, "sessions": sessions, "movements_count": len(movements)},
        "bank": {"credit": bank_credit, "debit": bank_debit, "net": round(bank_credit - bank_debit, 2), "entries": bank},
        "top_products": top_products,
    }


@router.get("/export/reports/daily.xlsx")
async def export_daily(date: Optional[str] = None, user: dict = Depends(require_module("reports"))):
    r = await daily_report(date, user)
    settings = await load_settings()
    method_labels = {"especes": "Espèces", "wave": "WAVE", "orange_money": "Orange Money", "cheque": "Chèque", "banque": "Banque"}
    rows = [["VENTES", "", "", ""], ["Nombre de factures", r["sales"]["count"], "", ""], ["Chiffre d'affaires", r["sales"]["total_revenue"], "", ""], ["Remises", r["sales"]["total_discount"], "", ""], ["Encaissé à la vente", r["sales"]["total_received"], "", ""], ["Vendu à crédit", r["sales"]["total_credit"], "", ""], ["Factures annulées", r["sales"]["cancelled_count"], "", ""], []]
    rows.append(["PAR MODE DE PAIEMENT", "Ventes", "Règlements", "Décaissements"])
    for m, v in r["by_method"].items():
        rows.append([method_labels.get(m, m), v["sales_cash"], round(v["invoice_payments"] + v["customer_payments"], 2), v["supplier_payments"]])
    rows += [[], ["CAISSE", "", "", ""], ["Entrées", r["cash"]["total_in"], "", ""], ["Sorties", r["cash"]["total_out"], "", ""], ["Net caisse", r["cash"]["net"], "", ""], [], ["BANQUE", "", "", ""], ["Crédits", r["bank"]["credit"], "", ""], ["Débits", r["bank"]["debit"], "", ""], ["Net banque", r["bank"]["net"], "", ""], [], ["FACTURES DU JOUR", "Client", "Total", "Reste dû"]]
    for s in r["sales"]["list"]:
        rows.append([s["invoice_number"] + (" (annulée)" if s["status"] == "cancelled" else ""), s["customer_name"], s["total"], s["balance_due"]])
    return build_workbook("Rapport journalier", ["Rubrique", "Valeur", "", ""], rows, settings["name"], f"Journée du {r['date']}")


HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def build_workbook(title: str, headers: list, rows: list, company: str, subtitle: str = "") -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([company])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([title + (f" — {subtitle}" if subtitle else "")])
    ws["A2"].font = Font(bold=True, size=11)
    ws.append([f"Édité le {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])
    ws.append(headers)
    for cell in ws[5]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for idx, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[idx - 1])) for r in rows if idx - 1 < len(r) and r[idx - 1] is not None] + [8])
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 60)
    ws.freeze_panes = "A6"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{title.lower().replace(' ', '_')}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def fmt_dt(v) -> str:
    try:
        d = v if isinstance(v, datetime) else datetime.fromisoformat(v)
        return d.strftime("%d/%m/%Y %H:%M")
    except Exception:
        return str(v or "")


MODULE_LABELS = {"pos": "Facturation", "sales": "Ventes", "products": "Produits", "categories": "Familles", "purchases": "Achats", "customers": "Clients", "suppliers": "Fournisseurs", "accounts": "Comptes tiers", "cash": "Gestion caisse", "bank": "Journal Banque", "depots": "Dépôts", "reports": "Rapports", "users": "Utilisateurs", "audit": "Historique", "settings": "Paramètres"}
ACTION_LABELS = {"create": "Création", "update": "Modification", "delete": "Suppression", "cancel": "Annulation", "return": "Retour"}
CATEGORY_LABELS = {"vente": "Vente", "reglement": "Règlement", "depense": "Dépense", "versement_banque": "Versement banque", "apport": "Apport", "autre": "Autre"}


def _diff(before, after) -> str:
    before, after = before or {}, after or {}
    keys = [k for k in dict.fromkeys(list(before) + list(after)) if k not in ("created_at", "updated_at", "id")]
    parts = [f"{k}: {before.get(k, '—')} → {after.get(k, '—')}" for k in keys if before.get(k) != after.get(k)]
    return "; ".join(str(p) for p in parts)[:1000]


@router.get("/export/audit-logs.xlsx")
async def export_audit(user_id: Optional[str] = None, module: Optional[str] = None, action: Optional[str] = None, start_date: Optional[str] = None, end_date: Optional[str] = None, user: dict = Depends(require_module("audit"))):
    query = {}
    if user_id:
        query["user_id"] = user_id
    if module:
        query["module"] = module
    if action:
        query["action"] = action
    if start_date or end_date:
        query["created_at"] = {}
        if start_date:
            query["created_at"]["$gte"] = start_date
        if end_date:
            query["created_at"]["$lte"] = end_date + "T23:59:59.999999" if len(end_date) == 10 else end_date
    logs = await db.audit_logs.find(query, {"_id": 0}).sort("created_at", -1).to_list(5000)
    settings = await load_settings()
    rows = [[fmt_dt(l["created_at"]), l.get("user_name"), l.get("user_role"), ACTION_LABELS.get(l["action"], l["action"]), MODULE_LABELS.get(l["module"], l["module"]), l.get("label") or l.get("entity_id"), _diff(l.get("before"), l.get("after"))] for l in logs]
    return build_workbook("Historique", ["Date", "Utilisateur", "Rôle", "Action", "Module", "Élément", "Modifications"], rows, settings["name"])


@router.get("/export/accounts/{entity_type}/{entity_id}/statement.xlsx")
async def export_statement(entity_type: str, entity_id: str, user: dict = Depends(require_module("accounts"))):
    data = await account_statement(entity_type, entity_id, user)
    settings = await load_settings()
    rows = [[fmt_dt(l["date"])[:10], l["label"], l.get("reference", ""), l["debit"] or None, l["credit"] or None, l["balance"]] for l in data["lines"]]
    rows.append(["", "TOTAUX", "", data["totals"]["debit"], data["totals"]["credit"], data["totals"]["balance"]])
    label = "Client" if entity_type == "customer" else "Fournisseur"
    return build_workbook("Releve de compte", ["Date", "Libellé", "Référence", "Facturé", "Réglé", "Solde"], rows, settings["name"], f"{label} {data['entity']['name']}")


@router.get("/export/cash/movements.xlsx")
async def export_cash(session_id: Optional[str] = None, date: Optional[str] = None, user: dict = Depends(require_module("cash"))):
    query = {}
    subtitle = ""
    if session_id:
        query["session_id"] = session_id
        session = await db.cash_sessions.find_one({"id": session_id}, {"_id": 0})
        if session:
            subtitle = f"Session du {fmt_dt(session['opened_at'])}"
    elif date:
        query["created_at"] = {"$gte": f"{date}T00:00:00", "$lte": f"{date}T23:59:59.999999"}
        subtitle = f"Journée du {date}"
    else:
        session = await db.cash_sessions.find_one({"status": "open"}, {"_id": 0})
        if session:
            query["session_id"] = session["id"]
            subtitle = f"Session en cours (ouverte le {fmt_dt(session['opened_at'])})"
        else:
            today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
            query["created_at"] = {"$gte": f"{today}T00:00:00", "$lte": f"{today}T23:59:59.999999"}
            subtitle = f"Journée du {today}"
    movements = await db.cash_movements.find(query, {"_id": 0}).sort("created_at", 1).to_list(10000)
    settings = await load_settings()
    rows = [[fmt_dt(m["created_at"]), m["description"], CATEGORY_LABELS.get(m.get("category"), m.get("category")), m.get("reference", ""), m.get("created_by"), m["amount"] if m["type"] == "in" else None, m["amount"] if m["type"] == "out" else None] for m in movements]
    total_in = sum(m["amount"] for m in movements if m["type"] == "in")
    total_out = sum(m["amount"] for m in movements if m["type"] == "out")
    rows.append(["", "TOTAUX", "", "", "", round(total_in, 2), round(total_out, 2)])
    rows.append(["", "NET", "", "", "", round(total_in - total_out, 2), None])
    return build_workbook("Journal de caisse", ["Date", "Description", "Catégorie", "Référence", "Par", "Entrée", "Sortie"], rows, settings["name"], subtitle)
