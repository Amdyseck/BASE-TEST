"""Iteration 5: daily report, non_paye payment method, xlsx export."""
import os
import io
import pytest
import requests
from datetime import datetime, timezone
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{API}/auth/login", json={"email": "admin@wakeur.sn", "password": "admin123"})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def cashier_token():
    r = requests.post(f"{API}/auth/login", json={"email": "caissier.test@wakeur.sn", "password": "test1234"})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def h_admin(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


@pytest.fixture(scope="module")
def h_cashier(cashier_token):
    return {"Authorization": f"Bearer {cashier_token}"}


@pytest.fixture(scope="module")
def a_product(h_admin):
    r = requests.get(f"{API}/products", headers=h_admin)
    assert r.status_code == 200
    prods = [p for p in r.json() if p.get("stock", 0) > 10]
    assert prods, "Need at least one product with stock>10"
    return prods[0]


@pytest.fixture(scope="module")
def a_customer(h_admin):
    r = requests.get(f"{API}/customers", headers=h_admin)
    assert r.status_code == 200
    cs = [c for c in r.json() if c.get("name") != "Client Comptoir"]
    assert cs, "Need at least one real customer"
    return cs[0]


# ---------- Daily Report ----------

class TestDailyReport:
    def test_daily_report_admin(self, h_admin):
        r = requests.get(f"{API}/reports/daily", headers=h_admin)
        assert r.status_code == 200, r.text
        d = r.json()
        for k in ["date", "sales", "by_method", "encaissements", "cash", "bank", "top_products"]:
            assert k in d
        for k in ["count", "total_revenue", "total_received", "total_credit", "total_discount", "cancelled_count", "list"]:
            assert k in d["sales"]
        for k in ["total_in", "total_out", "net", "by_category", "sessions", "movements_count"]:
            assert k in d["cash"]
        for k in ["credit", "debit", "net", "entries"]:
            assert k in d["bank"]

    def test_daily_report_with_date(self, h_admin):
        r = requests.get(f"{API}/reports/daily?date=2026-09-25", headers=h_admin)
        assert r.status_code == 200
        d = r.json()
        assert d["date"] == "2026-09-25"

    def test_daily_report_forbidden_cashier(self, h_cashier):
        r = requests.get(f"{API}/reports/daily", headers=h_cashier)
        assert r.status_code == 403

    def test_daily_report_coherence_cash_sale(self, h_admin, a_product):
        # Ensure a cash session is open (open one if not)
        sess = requests.get(f"{API}/cash/session/current", headers=h_admin)
        opened_here = False
        if sess.status_code != 200 or not sess.json():
            r = requests.post(f"{API}/cash/session/open", headers=h_admin, json={"opening_amount": 0})
            if r.status_code not in (200, 201):
                pytest.skip(f"Cannot open cash session: {r.status_code} {r.text}")
            opened_here = True

        before = requests.get(f"{API}/reports/daily", headers=h_admin).json()
        b_sales = before["sales"]["count"]
        b_cash_in = before["cash"]["total_in"]
        b_esp = before["by_method"].get("especes", {}).get("sales_cash", 0)

        unit_price = a_product["price"]
        payload = {
            "items": [{"product_id": a_product["id"], "product_name": a_product["name"], "quantity": 1, "unit_price": unit_price, "total": unit_price}],
            "discount": 0,
            "amount_paid": unit_price,
            "payment_method": "especes",
        }
        s = requests.post(f"{API}/sales", headers=h_admin, json=payload)
        assert s.status_code == 200, s.text

        after = requests.get(f"{API}/reports/daily", headers=h_admin).json()
        assert after["sales"]["count"] == b_sales + 1
        assert round(after["by_method"]["especes"]["sales_cash"] - b_esp, 2) >= round(unit_price - 0.01, 2)
        assert round(after["cash"]["total_in"] - b_cash_in, 2) >= round(unit_price - 0.01, 2)

    def test_export_daily_xlsx(self, h_admin):
        r = requests.get(f"{API}/export/reports/daily.xlsx", headers=h_admin)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # cells should contain "VENTES"
        found = any(cell.value == "VENTES" for row in ws.iter_rows() for cell in row)
        assert found

    def test_export_daily_xlsx_cashier_forbidden(self, h_cashier):
        r = requests.get(f"{API}/export/reports/daily.xlsx", headers=h_cashier)
        assert r.status_code == 403


# ---------- non_paye payment method ----------

class TestNonPayeSale:
    def test_non_paye_without_customer_fails(self, h_admin, a_product):
        unit = a_product["price"]
        payload = {
            "items": [{"product_id": a_product["id"], "product_name": a_product["name"], "quantity": 1, "unit_price": unit, "total": unit}],
            "discount": 0,
            "amount_paid": unit,
            "payment_method": "non_paye",
        }
        r = requests.post(f"{API}/sales", headers=h_admin, json=payload)
        assert r.status_code == 400
        assert "client" in r.json().get("detail", "").lower()

    def test_non_paye_with_customer_ok(self, h_admin, a_product, a_customer):
        # customer balance before
        cust_before = requests.get(f"{API}/customers", headers=h_admin).json()
        b_bal = next((c["balance"] for c in cust_before if c["id"] == a_customer["id"]), 0)

        unit = a_product["price"]
        payload = {
            "customer_id": a_customer["id"],
            # customer_name intentionally omitted
            "items": [{"product_id": a_product["id"], "product_name": a_product["name"], "quantity": 1, "unit_price": unit, "total": unit}],
            "discount": 0,
            "amount_paid": unit,  # will be forced to 0
            "payment_method": "non_paye",
        }
        r = requests.post(f"{API}/sales", headers=h_admin, json=payload)
        assert r.status_code == 200, r.text
        sale = r.json()
        assert sale["amount_paid"] == 0
        assert sale["balance_due"] == round(unit, 2)
        assert sale["is_paid"] is False
        assert sale["customer_name"] == a_customer["name"]

        # customer balance increased
        cust_after = requests.get(f"{API}/customers", headers=h_admin).json()
        a_bal = next((c["balance"] for c in cust_after if c["id"] == a_customer["id"]), 0)
        assert round(a_bal - b_bal, 2) == round(unit, 2)

        # sale visible in /sales with balance_due>0
        sales = requests.get(f"{API}/sales", headers=h_admin).json()
        found = next((s for s in sales if s["id"] == sale["id"]), None)
        assert found is not None
        assert found["balance_due"] > 0
