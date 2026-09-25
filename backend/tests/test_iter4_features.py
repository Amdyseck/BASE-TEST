"""
Iteration 4 backend tests:
1. PUT /api/auth/change-password (vendeur.test) - wrong current, correct, new-login, restore
2. GET /api/cash/sessions/{id} (session+movements+summary; 404 unknown)
3. Full cash flow (open->manual out->admin sale espèces auto in->close->detail summary)
4. GET /api/dashboard/overdue (creates a credit sale, days=0 includes, days=7 excludes)
5. Excel exports xlsx: audit-logs, statement, cash movements (permissions + openpyxl)
"""
import io
import os
import time
import pytest
import requests
from openpyxl import load_workbook

BASE = (os.environ.get("REACT_APP_BACKEND_URL") or open("/app/frontend/.env").read().split("REACT_APP_BACKEND_URL=")[1].split("\n")[0]).rstrip("/") + "/api"

ADMIN = ("admin@wakeur.sn", "admin123")
VENDEUR = ("vendeur.test@wakeur.sn", "test1234")
CAISSIER = ("caissier.test@wakeur.sn", "test1234")

STATE = {}


def login(email, password):
    r = requests.post(f"{BASE}/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, f"login {email} failed: {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="session")
def admin_h():
    return {"Authorization": f"Bearer {login(*ADMIN)}"}


@pytest.fixture(scope="session")
def caissier_h():
    return {"Authorization": f"Bearer {login(*CAISSIER)}"}


# =============== 1. CHANGE PASSWORD (vendeur.test) ===============

class TestChangePassword:
    def test_wrong_current_password(self):
        tok = login(*VENDEUR)
        h = {"Authorization": f"Bearer {tok}"}
        r = requests.put(f"{BASE}/auth/change-password",
                         json={"current_password": "wrong_pwd", "new_password": "newpass123"},
                         headers=h)
        assert r.status_code == 400, r.text

    def test_correct_change_then_login_then_restore(self):
        tok = login(*VENDEUR)
        h = {"Authorization": f"Bearer {tok}"}
        # change to new
        r = requests.put(f"{BASE}/auth/change-password",
                         json={"current_password": "test1234", "new_password": "brandnew123"},
                         headers=h)
        assert r.status_code == 200, r.text
        # login with new password
        tok2 = login("vendeur.test@wakeur.sn", "brandnew123")
        h2 = {"Authorization": f"Bearer {tok2}"}
        # restore original
        r = requests.put(f"{BASE}/auth/change-password",
                         json={"current_password": "brandnew123", "new_password": "test1234"},
                         headers=h2)
        assert r.status_code == 200
        # verify restored
        login("vendeur.test@wakeur.sn", "test1234")

    def test_audit_entry_created(self, admin_h):
        # Look for recent users update audit entry
        r = requests.get(f"{BASE}/audit-logs", headers=admin_h, params={"module": "users", "action": "update"})
        assert r.status_code == 200
        logs = r.json() if isinstance(r.json(), list) else r.json().get("logs", r.json())
        # Just check there is at least one recent update on users module
        assert isinstance(logs, list) and len(logs) > 0


# =============== 2 + 3. CASH FLOW + SESSION DETAIL ===============

class TestCashFlow:
    def test_close_any_open_session(self, admin_h):
        r = requests.get(f"{BASE}/cash/session/current", headers=admin_h)
        assert r.status_code == 200
        cur = r.json().get("session")
        if cur:
            expected = r.json().get("summary", {}).get("expected", cur.get("opening_amount", 0))
            requests.post(f"{BASE}/cash/session/close",
                          json={"closing_amount": expected, "notes": "prep"}, headers=admin_h)

    def test_open_session(self, admin_h):
        r = requests.post(f"{BASE}/cash/session/open",
                          json={"opening_amount": 10000, "notes": "TEST iter4"}, headers=admin_h)
        assert r.status_code == 200, r.text
        STATE["session_id"] = r.json()["id"]

    def test_manual_out(self, admin_h):
        r = requests.post(f"{BASE}/cash/movements",
                          json={"type": "out", "amount": 2000,
                                "description": "TEST depense iter4",
                                "category": "depense"}, headers=admin_h)
        assert r.status_code == 200, r.text

    def test_admin_sale_espece_auto_in(self, admin_h):
        prods = requests.get(f"{BASE}/products", headers=admin_h).json()
        p = next((x for x in prods if x.get("stock", 0) >= 1 and x.get("price", 0) > 0), None)
        if not p:
            pytest.skip("no product")
        it = {"product_id": p["id"], "product_name": p["name"], "quantity": 1,
              "unit_price": p["price"], "total": p["price"]}
        payload = {"items": [it], "subtotal": p["price"], "discount": 0, "total": p["price"],
                   "amount_paid": p["price"], "payment_method": "especes", "customer_id": None}
        r = requests.post(f"{BASE}/sales", json=payload, headers=admin_h)
        assert r.status_code == 200, r.text
        STATE["sale_price"] = p["price"]

    def test_session_detail_404(self, admin_h):
        r = requests.get(f"{BASE}/cash/sessions/does-not-exist", headers=admin_h)
        assert r.status_code == 404

    def test_session_detail_caissier_allowed(self, caissier_h):
        # caissier has module 'cash'
        r = requests.get(f"{BASE}/cash/sessions/{STATE['session_id']}", headers=caissier_h)
        assert r.status_code == 200

    def test_close_and_detail(self, admin_h):
        cur = requests.get(f"{BASE}/cash/session/current", headers=admin_h).json()
        expected = cur["summary"]["expected"]
        r = requests.post(f"{BASE}/cash/session/close",
                          json={"closing_amount": expected, "notes": "TEST close iter4"},
                          headers=admin_h)
        assert r.status_code == 200
        assert r.json()["difference"] == 0

        # Get session detail
        r = requests.get(f"{BASE}/cash/sessions/{STATE['session_id']}", headers=admin_h)
        assert r.status_code == 200, r.text
        d = r.json()
        assert "session" in d and "movements" in d and "summary" in d
        s = d["summary"]
        assert "total_in" in s and "total_out" in s and "expected" in s and "by_category" in s
        assert s["total_out"] >= 2000
        assert s["total_in"] >= STATE.get("sale_price", 0)
        # expected computed vs opening+in-out
        opening = d["session"]["opening_amount"]
        assert abs(s["expected"] - (opening + s["total_in"] - s["total_out"])) < 0.01


# =============== 4. OVERDUE ===============

class TestOverdue:
    def test_create_credit_sale_customer(self, admin_h):
        r = requests.post(f"{BASE}/customers",
                          json={"name": "TEST_Overdue", "phone": "+221771112233",
                                "address": "", "email": ""}, headers=admin_h)
        assert r.status_code == 200
        STATE["cust_id"] = r.json()["id"]

        # Need a cash session open to make a sale? sale with amount_paid=0 might not need cash movement
        prods = requests.get(f"{BASE}/products", headers=admin_h).json()
        p = next((x for x in prods if x.get("stock", 0) >= 1 and x.get("price", 0) > 0), None)
        if not p:
            pytest.skip("no product")
        it = {"product_id": p["id"], "product_name": p["name"], "quantity": 1,
              "unit_price": p["price"], "total": p["price"]}
        payload = {"items": [it], "subtotal": p["price"], "discount": 0, "total": p["price"],
                   "amount_paid": 0, "payment_method": "especes",
                   "customer_id": STATE["cust_id"]}
        r = requests.post(f"{BASE}/sales", json=payload, headers=admin_h)
        assert r.status_code == 200, r.text
        STATE["overdue_sale_id"] = r.json()["id"]

    def test_overdue_days0_includes(self, admin_h):
        r = requests.get(f"{BASE}/dashboard/overdue?days=0", headers=admin_h)
        assert r.status_code == 200, r.text
        d = r.json()
        assert "customers" in d and "total_overdue" in d
        mine = [c for c in d["customers"] if c["customer_id"] == STATE["cust_id"]]
        assert mine, f"TEST customer not in overdue: {d}"
        m = mine[0]
        assert m["overdue_amount"] > 0
        assert "phone" in m
        assert m["phone"] == "+221771112233"
        assert isinstance(m.get("invoices"), list) and len(m["invoices"]) >= 1
        assert "days_late" in m
        # total_overdue is sum of overdue_amount
        total = sum(c["overdue_amount"] for c in d["customers"])
        assert abs(total - d["total_overdue"]) < 0.01

    def test_overdue_days7_excludes_recent(self, admin_h):
        r = requests.get(f"{BASE}/dashboard/overdue?days=7", headers=admin_h)
        assert r.status_code == 200
        d = r.json()
        mine = [c for c in d["customers"] if c["customer_id"] == STATE["cust_id"]]
        assert not mine, "Recent sale should be excluded from days=7"

    def test_cleanup_overdue(self, admin_h):
        if STATE.get("overdue_sale_id"):
            # cancel sale first to clear balance
            requests.post(f"{BASE}/sales/{STATE['overdue_sale_id']}/cancel", headers=admin_h)
        if STATE.get("cust_id"):
            requests.delete(f"{BASE}/customers/{STATE['cust_id']}", headers=admin_h)


# =============== 5. EXCEL EXPORTS ===============

class TestExports:
    def _assert_xlsx(self, r):
        assert r.status_code == 200, r.text
        ct = r.headers.get("Content-Type", "")
        assert "spreadsheetml" in ct or "xlsx" in ct, ct
        wb = load_workbook(io.BytesIO(r.content))
        return wb

    def test_audit_export_admin(self, admin_h):
        r = requests.get(f"{BASE}/export/audit-logs.xlsx", headers=admin_h)
        wb = self._assert_xlsx(r)
        ws = wb.active
        headers = [ws.cell(row=5, column=i).value for i in range(1, 8)]
        assert headers == ["Date", "Utilisateur", "Rôle", "Action", "Module", "Élément", "Modifications"], headers

    def test_audit_export_filter_module(self, admin_h):
        r = requests.get(f"{BASE}/export/audit-logs.xlsx?module=users", headers=admin_h)
        self._assert_xlsx(r)

    def test_audit_export_caissier_403(self, caissier_h):
        r = requests.get(f"{BASE}/export/audit-logs.xlsx", headers=caissier_h)
        assert r.status_code == 403

    def test_statement_export(self, admin_h):
        # pick any customer
        custs = requests.get(f"{BASE}/customers", headers=admin_h).json()
        if not custs:
            pytest.skip("no customer")
        cid = custs[0]["id"]
        r = requests.get(f"{BASE}/export/accounts/customer/{cid}/statement.xlsx", headers=admin_h)
        wb = self._assert_xlsx(r)
        ws = wb.active
        # last row should have TOTAUX label in column B
        last_row = ws.max_row
        cell_b = ws.cell(row=last_row, column=2).value
        assert cell_b == "TOTAUX", f"expected TOTAUX in last row, got {cell_b}"

    def test_cash_export_admin(self, admin_h):
        r = requests.get(f"{BASE}/export/cash/movements.xlsx", headers=admin_h)
        self._assert_xlsx(r)

    def test_cash_export_with_session(self, admin_h):
        sid = STATE.get("session_id")
        if not sid:
            pytest.skip("no session")
        r = requests.get(f"{BASE}/export/cash/movements.xlsx?session_id={sid}", headers=admin_h)
        self._assert_xlsx(r)

    def test_cash_export_vendeur_403(self):
        tok = login(*VENDEUR)
        h = {"Authorization": f"Bearer {tok}"}
        r = requests.get(f"{BASE}/export/cash/movements.xlsx", headers=h)
        assert r.status_code == 403, r.text
